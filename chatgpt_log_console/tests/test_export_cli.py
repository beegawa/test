"""더블클릭 한 번으로 대화를 엑셀로 뽑는 도구 (export_conversations.py)."""

from io import BytesIO

import pytest
from conftest import jsonl

import export_conversations as 내보내기
import keystore
import settings
from compliance import ComplianceClient
from records import normalize


def 대화(번호, 역할, 글, 대화id, 시각):
    return {
        "event_id": f"cm-{번호}", "type": "CONVERSATION_MESSAGE",
        "actor": {"type": "ACCOUNT_USER", "user_id": "u1", "user_email": "hong@example.com"},
        "timestamp": 시각,
        "message": {"id": f"m{번호}", "author": {"type": 역할},
                    "content": {"type": "text", "value": 글}},
        "conversation": {"id": 대화id, "title": "New chat"},
    }


@pytest.fixture
def 환경(tmp_path, monkeypatch, session):
    monkeypatch.setenv("CHATGPT_LOG_NO_KEYRING", "1")
    monkeypatch.delenv("CHATGPT_ADMIN_KEY", raising=False)
    monkeypatch.setattr(settings, "DATA_DIR", tmp_path)
    monkeypatch.setattr(keystore, "DATA_DIR", tmp_path)
    monkeypatch.setattr(
        내보내기, "ComplianceClient",
        lambda key: ComplianceClient(key, base_url="https://e.test", workspace_id="ws_1",
                                     org_id=None, session=session, sleep=lambda _s: None))
    session.event_types_ok = {"CONVERSATION_MESSAGE"}
    session.pages = [{"data": [{"id": "log_1"}], "has_more": False,
                      "last_end_time": "2026-09-17T00:00:00Z"}]
    session.logs["log_1"] = jsonl(
        대화(1, "user", "연차 규정 알려줘", "c-1", "2026-09-16T01:00:00Z"),
        대화(2, "assistant", "12조를 보세요.", "c-1", "2026-09-16T01:00:05Z"),
        대화(3, "user", "다른 질문", "c-2", "2026-09-16T02:00:00Z"),
    )
    return tmp_path


def 엑셀읽기(경로):
    from openpyxl import load_workbook
    return load_workbook(BytesIO(경로.read_bytes())).active


def 만들어진파일(폴더):
    파일들 = sorted(폴더.glob("*.xlsx"))
    assert 파일들, "엑셀 파일이 만들어지지 않았습니다"
    return 파일들[-1]


def test_키가_없으면_안내하고_종료코드_2(환경, tmp_path, capsys):
    assert 내보내기.main(["--out", str(tmp_path), "--no-open"]) == 2
    assert "관리자 키가 없습니다" in capsys.readouterr().out


def test_받아서_엑셀로_만든다(환경, tmp_path):
    keystore.save_key("sk-admin-1234567890")
    나가는곳 = tmp_path / "out"

    assert 내보내기.main(["--out", str(나가는곳), "--no-open", "--db", str(tmp_path / "a.db")]) == 0

    시트 = 엑셀읽기(만들어진파일(나가는곳))
    assert 시트.max_row == 4                      # 머리글 1 + 메시지 3
    assert 시트.cell(row=1, column=1).value == "시간(UTC)"


def test_같은_대화의_질문과_답변이_붙어_나온다(환경, tmp_path):
    keystore.save_key("sk-admin-1234567890")
    나가는곳 = tmp_path / "out"
    내보내기.main(["--out", str(나가는곳), "--no-open", "--db", str(tmp_path / "a.db")])

    시트 = 엑셀읽기(만들어진파일(나가는곳))
    대화열 = [시트.cell(row=r, column=5).value for r in range(2, 5)]
    요약열 = [시트.cell(row=r, column=6).value for r in range(2, 5)]
    assert 대화열 == ["c-1", "c-1", "c-2"]          # 같은 대화끼리 모여 있다
    assert 요약열[0].startswith("사용자: ")          # 그 안에서는 시간 순
    assert 요약열[1].startswith("ChatGPT: ")


def test_사용자와_키워드로_걸러_낸다(환경, tmp_path):
    keystore.save_key("sk-admin-1234567890")
    db = str(tmp_path / "a.db")
    내보내기.main(["--out", str(tmp_path / "all"), "--no-open", "--db", db])

    내보내기.main(["--skip-pull", "--q", "연차", "--out", str(tmp_path / "kw"),
                 "--no-open", "--db", db])
    assert 엑셀읽기(만들어진파일(tmp_path / "kw")).max_row == 2      # 머리글 + 1건

    assert 내보내기.main(["--skip-pull", "--user", "없는사람", "--out", str(tmp_path / "no"),
                       "--no-open", "--db", db]) == 1               # 결과 없음


def test_두_번_돌려도_중복되지_않는다(환경, tmp_path, session):
    keystore.save_key("sk-admin-1234567890")
    db = str(tmp_path / "a.db")
    내보내기.main(["--out", str(tmp_path / "a"), "--no-open", "--db", db])
    session.pages = [{"data": [{"id": "log_1"}], "has_more": False,
                      "last_end_time": "2026-09-17T00:00:00Z"}]
    내보내기.main(["--out", str(tmp_path / "b"), "--no-open", "--db", db])

    assert 엑셀읽기(만들어진파일(tmp_path / "b")).max_row == 4       # 그대로 3건


def test_엑셀_파일은_소유자만_읽는다(환경, tmp_path):
    import os
    import stat

    keystore.save_key("sk-admin-1234567890")
    나가는곳 = tmp_path / "out"
    내보내기.main(["--out", str(나가는곳), "--no-open", "--db", str(tmp_path / "a.db")])
    모드 = stat.S_IMODE(os.stat(만들어진파일(나가는곳)).st_mode)
    assert oct(모드) == "0o600"


def test_대화순_정렬은_대화id_없는_것도_떨어뜨리지_않는다():
    행들 = [{"conversation_id": "", "ts": "2026-09-16T00:00:00Z", "id": "x"},
           {"conversation_id": "c-1", "ts": "2026-09-16T00:00:00Z", "id": "y"}]
    assert len(내보내기.대화순_정렬(행들)) == 2


def test_진행률은_같은_줄을_덮어쓰고_너무_자주_찍지_않는다(capsys):
    표시 = 내보내기.진행표시(간격=999)         # 간격이 길면 첫 번째만 찍힌다
    표시({"listed": 10, "fetched": 5, "saved": 3})
    표시({"listed": 20, "fetched": 15, "saved": 9})
    출력 = capsys.readouterr().out
    assert 출력.startswith("\r")
    assert "목록 10건" in 출력
    assert "목록 20건" not in 출력             # 간격 안에 들어온 것은 건너뛴다

    표시.끝()
    assert "\r" in capsys.readouterr().out     # 줄을 지운다


def test_진행률의_경과_시간_표시():
    시간 = 내보내기.진행표시._시간
    assert 시간(45) == "45초"
    assert 시간(125) == "2분 5초"
    assert 시간(3700) == "1시간 1분"


def test_중간에_끊어도_받은_것은_남는다(환경, tmp_path, monkeypatch):
    """Ctrl+C 로 멈춰도 그때까지 저장된 것은 지키고, 다시 실행하면 이어 받는다."""
    keystore.save_key("sk-admin-1234567890")
    db = str(tmp_path / "a.db")

    def 중단(*_a, **_k):
        raise KeyboardInterrupt

    monkeypatch.setattr(내보내기, "collect", 중단)
    assert 내보내기.main(["--out", str(tmp_path / "x"), "--no-open", "--db", db]) == 130
